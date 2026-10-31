"""Text extraction that separates visible text from hidden text.

Hidden text is a classic way to smuggle instructions to an AI screener:
white text, microscopic fonts, text placed off the page, or Word's "hidden" style.
It is removed from what the agents read, and kept aside so the Guard can scan
and report it.

Known limitation: text covered by a drawn shape or image is not detected.
"""
import io
from dataclasses import dataclass, field

import docx
import pymupdf

_NEAR_WHITE = 235  # every RGB channel at or above this counts as white-on-white
_MIN_FONT_PT = 4.0


@dataclass
class HiddenSegment:
    text: str
    reason: str  # white_text | tiny_font | off_page | invisible | docx_hidden


@dataclass
class Extraction:
    visible_text: str
    hidden: list[HiddenSegment] = field(default_factory=list)


def extract(content: bytes, kind: str) -> Extraction:
    if kind == "pdf":
        return _extract_pdf(content)
    if kind == "docx":
        return _extract_docx(content)
    raise ValueError(f"unsupported kind {kind!r}")


def _is_near_white(rgb: tuple[int, int, int]) -> bool:
    return min(rgb) >= _NEAR_WHITE


def _pdf_hidden_reason(span: dict, page_rect: pymupdf.Rect) -> str | None:
    color = int(span.get("color", 0))
    if _is_near_white(((color >> 16) & 255, (color >> 8) & 255, color & 255)):
        return "white_text"
    if float(span.get("size", 12)) < _MIN_FONT_PT:
        return "tiny_font"
    if span.get("alpha", 255) == 0:
        return "invisible"
    bbox = pymupdf.Rect(span["bbox"])
    area = bbox.get_area()
    if area > 0 and (bbox & page_rect).get_area() < 0.5 * area:
        return "off_page"
    return None


def _extract_pdf(content: bytes) -> Extraction:
    lines: list[str] = []
    hidden: list[HiddenSegment] = []
    # Keep text outside the page box too, so off-page smuggling is detected and reported.
    text_flags = pymupdf.TEXTFLAGS_DICT & ~pymupdf.TEXT_MEDIABOX_CLIP
    with pymupdf.open(stream=content, filetype="pdf") as doc:
        for page in doc:
            page_rect = page.rect
            for block in page.get_text("dict", flags=text_flags, clip=pymupdf.INFINITE_RECT())["blocks"]:
                if block.get("type") != 0:
                    continue
                for line in block["lines"]:
                    parts: list[str] = []
                    for span in line["spans"]:
                        text = span["text"]
                        if not text.strip():
                            parts.append(text)
                            continue
                        reason = _pdf_hidden_reason(span, page_rect)
                        if reason:
                            hidden.append(HiddenSegment(text.strip(), reason))
                        else:
                            parts.append(text)
                    line_text = "".join(parts).strip()
                    if line_text:
                        lines.append(line_text)
                lines.append("")
    return Extraction("\n".join(lines), hidden)


def _docx_hidden_reason(run) -> str | None:
    font = run.font
    if font.hidden:
        return "docx_hidden"
    try:
        rgb = font.color.rgb if font.color is not None and font.color.type is not None else None
    except (AttributeError, ValueError):
        rgb = None
    if rgb is not None and _is_near_white((rgb[0], rgb[1], rgb[2])):
        return "white_text"
    if font.size is not None and font.size.pt < _MIN_FONT_PT:
        return "tiny_font"
    return None


def _docx_paragraphs(document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in document.sections:
        for part in (section.header, section.footer):
            yield from part.paragraphs


def _extract_docx(content: bytes) -> Extraction:
    document = docx.Document(io.BytesIO(content))
    lines: list[str] = []
    hidden: list[HiddenSegment] = []
    for paragraph in _docx_paragraphs(document):
        parts: list[str] = []
        for run in paragraph.runs:
            if not run.text.strip():
                parts.append(run.text)
                continue
            reason = _docx_hidden_reason(run)
            if reason:
                hidden.append(HiddenSegment(run.text.strip(), reason))
            else:
                parts.append(run.text)
        text = "".join(parts).strip()
        if text:
            lines.append(text)
    return Extraction("\n".join(lines), hidden)
